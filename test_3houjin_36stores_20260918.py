# -*- coding: utf-8 -*-
"""
3法人36店対応（2026-09-18）の検証テスト。

【何をするスクリプトか】
本番の Google スプレッドシート・メール・ブラウザには一切つながず、ダミー（偽物）と純関数だけで
3法人対応の追加ぶんを確かめる。確認する対象：
  1. 店舗マスタ（stores_config）… 36店・法人別の件数・重複なし・飛鳥22店が事務ポータルの店名と一致・シート名安全
  2. 合言葉の判定（app_logic.evaluate_password）… 一致・不一致・重複・本部用・後方互換
     ＋2026-09-19 名前ごとの管理者7人（[admin_passwords]）… 7人それぞれ／旧 admin_password／両方／重複拒否／空 dict
  3. 受取時期（app_logic.pickup_options / pickup_cap / resolve_pickup / _pickup_display）… 便の月・丸め・旧予約の互換
  4. 3,000円ライン（app_logic.box_totals）… 確定版 表の例A（3,600円到達）・例B（2,000円未達）
  5. 引取依頼書（app_logic.build_pickup_request → yuzu_core.write_pickup_request_excel）… 11列・式・記名欄・合計・同一法人注記・幅
  6. メール（mailer）… smtplib を偽物に差し替え、接続1回・ログイン1回・35通・リンク・宛先未登録の要約・途中失敗・予算切れ
  7. Googleシート（gsheet_store）… values_batch_get 1回／失敗時36回／_店舗情報 が無いブック
  8. 見本・生成スクリプト・画面（2026-09-19）… secrets.toml.sample が TOML として読め [admin_passwords] が7人／
     _検証用合言葉を作る.py の出力が TOML として読め 7＋36 で全値が別／AppTest（ダミー Secrets・ローカル保管庫）で
     名前つき管理者で入ると上の帯に名前が出て、法人で絞って36店を切り替えられる／左バーには何も無い
     （2026-09-19 左バー→上の帯。AppTest は at.main（本体）で探し、at.sidebar は「空であること」だけ見る）
  ★合言葉の値・メールアドレス・所在地・人名はすべてダミー（実在の値は一切書かない。管理者の「名前」は役職名の固定7つ）。

【実行方法（このフォルダで）】
    python test_3houjin_36stores_20260918.py
※ Google API は叩かない。本番シートには書き込まない。メールも実送信しない。追加ライブラリ不要（openpyxl は導入済み前提）。
"""

import io
import os
import re
import sys
import time

# このフォルダを import パスに入れる
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import stores_config as sc
import yuzu_core as yc
import app_logic as al
import gsheet_store as gs
import mailer

from openpyxl import load_workbook

_passed = 0
_failed = 0


def check(name, cond, detail=''):
    global _passed, _failed
    if cond:
        _passed += 1
        print('  [OK] ' + name)
    else:
        _failed += 1
        print('  [NG] ' + name + ('  … ' + str(detail) if detail else ''))


# ============================================================================
# 1) 店舗マスタ
# ============================================================================
print('\n■ 1) 店舗マスタ（36店）')
check('len(STORE_NAMES)==36', len(sc.STORE_NAMES) == 36, len(sc.STORE_NAMES))
check('SOYOUZ 9・NAIKANDO 5・ASUKA 22', (len(sc.SOYOUZ), len(sc.NAIKANDO), len(sc.ASUKA)) == (9, 5, 22))
check('重複0', len(set(sc.STORE_NAMES)) == 36)
check('STORE_COUNT==36', sc.STORE_COUNT == 36)
check('COMPANY_OF 36件・飛鳥22件', len(sc.COMPANY_OF) == 36 and sum(1 for v in sc.COMPANY_OF.values() if v == '飛鳥') == 22)
check('STORE_NAMES の並び＝SOYOUZ+NAIKANDO+ASUKA', sc.STORE_NAMES == sc.SOYOUZ + sc.NAIKANDO + sc.ASUKA)
check('36店すべて _safe_sheet_name(name)==name', all(yc._safe_sheet_name(n) == n for n in sc.STORE_NAMES))
# 事務ポータル飛鳥版の店名（api-client.js）と完全一致するか（ファイルがある環境でだけ検査）
_js = os.path.normpath(os.path.join(HERE, '..', '..', '..', '事務マニュアル', '事務ポータル', 'js', 'api-client.js'))
if os.path.exists(_js):
    names = []
    for ln in open(_js, encoding='utf-8'):
        m = re.search(r"会社: '飛鳥', 店名: '([^']+)'", ln)
        if m and '開発' not in ln:
            names.append(m.group(1))
    check('飛鳥22店が事務ポータル api-client.js と表記・順序とも一致', names == sc.ASUKA, names)
else:
    print('  [--] api-client.js が見つからないため店名照合は省略')

# ============================================================================
# 2) 合言葉の判定（値はダミー）
# ============================================================================
print('\n■ 2) 合言葉の判定')
SP = {n: 'dummy-%02d' % i for i, n in enumerate(sc.STORE_NAMES)}   # 36店に別々のダミー値
r = al.evaluate_password('dummy-03', SP, 'dummy-admin', '', sc.STORE_NAMES)
check('店の合言葉→その店に確定（東立石）', r['ok'] and r['mode'] == 'store' and r['store'] == '東立石', r)
r = al.evaluate_password('dummy-14', SP, 'dummy-admin', '', sc.STORE_NAMES)
check('店の合言葉→飛鳥 本店に確定', r['ok'] and r['store'] == '本店', r)
r = al.evaluate_password('dummy-admin', SP, 'dummy-admin', '', sc.STORE_NAMES)
check('本部用→admin モード（店は未確定＝自由切替）', r['ok'] and r['mode'] == 'admin' and r['store'] is None, r)
r = al.evaluate_password('wrong', SP, 'dummy-admin', '', sc.STORE_NAMES)
check('不一致→拒否', (not r['ok']) and '違います' in r['error'], r)
r = al.evaluate_password('', SP, 'dummy-admin', '', sc.STORE_NAMES)
check('空→拒否', not r['ok'], r)
SP2 = dict(SP); SP2['柏'] = SP2['東立石']       # 2店に同じ値
r = al.evaluate_password(SP2['東立石'], SP2, 'dummy-admin', '', sc.STORE_NAMES)
check('2店に同じ値→その値では入れず「設定を確認」', (not r['ok']) and '設定を確認' in r['error'], r)
SP3 = dict(SP); SP3['存在しない店'] = 'dummy-x'
r = al.evaluate_password('dummy-x', SP3, 'dummy-admin', '', sc.STORE_NAMES)
check('店舗マスタに無い店の合言葉→拒否＋設定確認', (not r['ok']) and '設定を確認' in r['error'], r)
r = al.evaluate_password('shared', None, '', 'shared', sc.STORE_NAMES)
check('後方互換：[store_passwords] 無し・app_password 一致→legacy（店を自由に選べる）', r['ok'] and r['mode'] == 'legacy', r)
r = al.evaluate_password('shared', SP, '', 'shared', sc.STORE_NAMES)
check('[store_passwords] があるとき app_password は無視される', not r['ok'], r)
r = al.evaluate_password('', None, '', '', sc.STORE_NAMES)
check('何も設定なし→dev（開発モード）', r['ok'] and r['mode'] == 'dev', r)

# --- 2026-09-19 名前ごとの管理者7人（[admin_passwords]）。値はダミー ---
print('\n■ 2b) 名前ごとの管理者（[admin_passwords]・7人）')
check('ADMIN_NAMES は固定の7人・この順', sc.ADMIN_NAMES == ['本間', '森田', '小島', 'ソユーズ担当者', '内観堂担当者', '加藤', '飛鳥薬局担当者'], sc.ADMIN_NAMES)
AP = {n: 'dummy-adm-%02d' % i for i, n in enumerate(sc.ADMIN_NAMES)}   # 7人に別々のダミー値
for n in sc.ADMIN_NAMES:
    r = al.evaluate_password(AP[n], SP, '', '', sc.STORE_NAMES, admin_passwords=AP)
    check('管理者 %s の値→admin・admin_name 一致・店は未確定' % n,
          r['ok'] and r['mode'] == 'admin' and r['admin_name'] == n and r['store'] is None, r)
r = al.evaluate_password('dummy-admin', SP, 'dummy-admin', '', sc.STORE_NAMES)
check('旧 admin_password だけ（admin_passwords 省略）→admin・admin_name=本部', r['ok'] and r['mode'] == 'admin' and r['admin_name'] == '本部', r)
r = al.evaluate_password('dummy-admin', SP, 'dummy-admin', '', sc.STORE_NAMES, admin_passwords=None)
check('旧 admin_password だけ（admin_passwords=None）→admin・本部', r['ok'] and r['mode'] == 'admin' and r['admin_name'] == '本部', r)
r1 = al.evaluate_password(AP['小島'], SP, 'dummy-admin', '', sc.STORE_NAMES, admin_passwords=AP)
r2 = al.evaluate_password('dummy-admin', SP, 'dummy-admin', '', sc.STORE_NAMES, admin_passwords=AP)
check('[admin_passwords] と旧 admin_password の両方あり→両方通る（小島／本部）',
      r1['ok'] and r1['admin_name'] == '小島' and r2['ok'] and r2['admin_name'] == '本部', (r1, r2))
r = al.evaluate_password('dummy-03', SP, 'dummy-admin', '', sc.STORE_NAMES, admin_passwords=AP)
check('店の合言葉は従来どおり店に固定・admin_name は None', r['ok'] and r['mode'] == 'store' and r['store'] == '東立石' and r['admin_name'] is None, r)
AP_S = dict(AP); AP_S['森田'] = SP['柏']            # 管理者と店で同じ値
r = al.evaluate_password(SP['柏'], SP, '', '', sc.STORE_NAMES, admin_passwords=AP_S)
check('管理者と店で同じ値→拒否＋警告文', (not r['ok']) and r['error'] == '同じ合言葉が複数に登録されています。管理本部に連絡し、設定を確認してください。', r)
AP_A = dict(AP); AP_A['加藤'] = AP_A['本間']         # 管理者どうし同じ値
r = al.evaluate_password(AP_A['本間'], SP, '', '', sc.STORE_NAMES, admin_passwords=AP_A)
check('管理者どうし同じ値→拒否＋警告文', (not r['ok']) and '複数に登録' in r['error'], r)
r = al.evaluate_password(AP['本間'], SP, AP['本間'], '', sc.STORE_NAMES, admin_passwords=AP)
check('名前つき管理者と旧 admin_password が同じ値→拒否', (not r['ok']) and '複数に登録' in r['error'], r)
r = al.evaluate_password(SP2['東立石'], SP2, '', '', sc.STORE_NAMES, admin_passwords=AP)
check('店どうし同じ値→拒否（文言は管理者と共通）', (not r['ok']) and '複数に登録' in r['error'], r)
r = al.evaluate_password('dummy-admin', SP, 'dummy-admin', '', sc.STORE_NAMES, admin_passwords={})
check('[admin_passwords] が空 dict→旧 admin_password にフォールバック（本部）', r['ok'] and r['mode'] == 'admin' and r['admin_name'] == '本部', r)
r = al.evaluate_password('dummy-03', SP, '', '', sc.STORE_NAMES, admin_passwords={})
check('[admin_passwords] が空 dict・旧も無し→店の合言葉は通る', r['ok'] and r['store'] == '東立石', r)
r = al.evaluate_password('wrong', SP, '', '', sc.STORE_NAMES, admin_passwords=AP)
check('管理者・店のどれとも違う→拒否', (not r['ok']) and '違います' in r['error'], r)
r = al.evaluate_password('', SP, '', '', sc.STORE_NAMES, admin_passwords=AP)
check('空欄→拒否（管理者の空値にも一致しない）', not r['ok'], r)
AP_E = dict(AP); AP_E['加藤'] = ''                  # まだ配らない人は空
r = al.evaluate_password('', None, '', '', sc.STORE_NAMES, admin_passwords={'加藤': ''})
check('管理者の値が全部空・他も無し→dev（未設定と同じ）', r['ok'] and r['mode'] == 'dev', r)
r = al.evaluate_password(AP['本間'], None, '', '', sc.STORE_NAMES, admin_passwords=AP_E)
check('[admin_passwords] だけ（店ごと・共有なし）でも管理者は入れる', r['ok'] and r['admin_name'] == '本間', r)
r = al.evaluate_password('wrong', None, '', '', sc.STORE_NAMES, admin_passwords=AP_E)
check('[admin_passwords] だけ・不一致→拒否（dev に落ちない）', (not r['ok']) and '違います' in r['error'], r)
check('戻り値のキーは常に ok/mode/store/admin_name/error',
      all(set(al.evaluate_password(p, s, a, l, sc.STORE_NAMES, admin_passwords=x).keys()) == {'ok', 'mode', 'store', 'admin_name', 'error'}
          for p, s, a, l, x in [('', None, '', '', None), ('x', SP, '', '', AP), ('shared', None, '', 'shared', None), ('dummy-admin', None, 'dummy-admin', '', None)]))

# ============================================================================
# 3) 受取時期（便の月）
# ============================================================================
print('\n■ 3) 受取時期＝便の月')
expect = {
    '202611': [(0, '今すぐ（随時便）'), (2, '2027年1月便（2ヶ月後）'), (5, '2027年4月便（5ヶ月後）')],
    '202612': [(0, '今すぐ（随時便）'), (1, '2027年1月便（1ヶ月後）'), (4, '2027年4月便（4ヶ月後）')],
    '202701': [(0, '今すぐ／2027年1月便（今月・18〜20日発送）'), (3, '2027年4月便（3ヶ月後）')],
    '202702': [(0, '今すぐ（随時便）'), (2, '2027年4月便（2ヶ月後）'), (5, '2027年7月便（5ヶ月後）')],
    '202709': [(0, '今すぐ（随時便）'), (1, '2027年10月便（1ヶ月後）'), (4, '2028年1月便（4ヶ月後）')],
    '202710': [(0, '今すぐ／2027年10月便（今月・18〜20日発送）'), (3, '2028年1月便（3ヶ月後）')],
}
for ym, exp in expect.items():
    got = [(o['offset'], o['label']) for o in al.pickup_options(ym)]
    check('pickup_options(%s)' % ym, got == exp, got)
check('選択肢の ym が YYYYMM（保存形式は不変）', all(len(o['ym']) == 6 and o['ym'].isdigit() for o in al.pickup_options('202611')))
check('発送日＝第3週の月〜水（付録B と一致：1/18-20, 4/19-21, 7/19-21, 10/18-20）',
      [yc.bin_ship_days(y) for y in ('202701', '202704', '202707', '202710')] == [(18, 20), (19, 21), (19, 21), (18, 20)])
check('丸め：当月202611・期限2027/02/28→4月便(5)を選ぶと1月便(2)', al.resolve_pickup(5, '2027/02/28', '202611') == (2, '202701'))
check('丸め：当月202611・期限2026/12/31→今すぐ(0)', al.resolve_pickup(5, '2026/12/31', '202611') == (0, '202611'))
check('期限なし→最も先の便まで選べる', al.pickup_cap('', '202611') == 5)
check('期限内なら選んだ便のまま（1月便・期限2027/03/31）', al.resolve_pickup(2, '2027/03/31', '202611') == (2, '202701'))
check('既存予約（202612・当月202611）の表示は 2026/12 を含む', '2026/12' in al._pickup_display('202612', '202611'), al._pickup_display('202612', '202611'))
check('便の月（202701・当月202611）の表示は 2027年1月便', al._pickup_display('202701', '202611') == '2027年1月便')
check('当月と同じ→今すぐ', al._pickup_display('202611', '202611') == '今すぐ')
check('pickup_label(2,"202701")==2027年1月便／pickup_label(1,"202612")==1ヶ月後／pickup_label(0)==今すぐ',
      (yc.pickup_label(2, '202701'), yc.pickup_label(1, '202612'), yc.pickup_label(0)) == ('2027年1月便', '1ヶ月後', '今すぐ'))
# 既存の予約行が有効のまま読めるか（reservation_map）
rm = al.reservation_map([{'予約した店': '東立石', '出し手店': '氷川台', '対象年月': '202611',
                          '受取予定月': '202612', '予約キー': 'K1', '薬品名': 'A', '予約日時': ''}], '202611')
check('既存予約（旧3ヶ月方式・202612）は当月202611で有効・ラベル 1ヶ月後', ('氷川台', 'K1') in rm and rm[('氷川台', 'K1')]['受取ラベル'] == '1ヶ月後', rm)

# ============================================================================
# 4) 3,000円ライン（確定版 表 例A・例B）
# ============================================================================
print('\n■ 4) box_totals（3,000円ライン）')
def _pr(store, key, amt, price=None, qty=1):
    return {'出し手店': store, '_ex_key': key, '薬品名': '薬' + key, '在庫金額': amt, '在庫数': qty,
            '薬価': (price if price is not None else amt / qty), '単位': '錠', '有効期限': '2027/06/30',
            '区分': '', '医薬品CD': key, 'ロットNO': 'L' + key}
result = {'proposal_rows': [
    _pr('B', 'b1', 1800), _pr('B', 'b2', 600), _pr('B', 'b3', 600), _pr('B', 'b4', 600),   # 例A 3,600円
    _pr('C', 'c1', 500), _pr('C', 'c2', 500), _pr('C', 'c3', 500), _pr('C', 'c4', 500),   # 例B 2,000円
]}
resv = [{'予約した店': 'A', '出し手店': s, '対象年月': '202611', '受取予定月': '202701', '予約キー': k, '薬品名': '薬' + k, '予約日時': 'now'}
        for s, k in [('B', 'b1'), ('B', 'b2'), ('B', 'b3'), ('B', 'b4'), ('C', 'c1'), ('C', 'c2'), ('C', 'c3'), ('C', 'c4')]]
bt = al.box_totals(result, 'A', resv, '202611')
recv = {b['相手店']: b for b in bt['receive']}
check('受け手A：Bは3,600円で到達', recv['B']['合計'] == 3600 and recv['B']['到達'] and recv['B']['件数'] == 4, recv.get('B'))
check('受け手A：Cは2,000円で未達（あと1,000円）', recv['C']['合計'] == 2000 and (not recv['C']['到達']) and recv['C']['あと'] == 1000, recv.get('C'))
check('受け手Aの supply は空（Aは出していない）', bt['supply'] == [])
btB = al.box_totals(result, 'B', resv, '202611')
check('出し手B：予約してきた店Aの合計3,600円・到達', btB['supply'] and btB['supply'][0]['相手店'] == 'A' and btB['supply'][0]['合計'] == 3600 and btB['supply'][0]['到達'])
check('box_min は CONFIG（3000）', bt['box_min'] == yc.CONFIG['box_min_amount'] == 3000)
plan = al.plan_reservations([], 'A', '202611', [{'_出し手店': 'C', '_key': 'c1', '薬品名': '薬c1', '_受取予定月': '202701', '_now': 'now'}])
check('3,000円未満でも予約は保存される（止めない）', plan['added'] == 1)
check('CONFIG: min_supply_amount=500／anytime_ship_amount=1500', yc.CONFIG['min_supply_amount'] == 500 and yc.CONFIG['anytime_ship_amount'] == 1500)

# ============================================================================
# 5) 引取依頼書（11列・式・記名欄・合計・同一法人注記・幅）
# ============================================================================
print('\n■ 5) 引取依頼書')
# 受け手＝東立石（ソユーズ）、出し手＝氷川台（内観堂＝別法人）と 柏（ソユーズ＝同一法人）
result2 = {'proposal_rows': [
    _pr('氷川台', 'h1', 1800.0, price=18.0, qty=100),   # 18円×100錠＝1,800円
    _pr('氷川台', 'h2', 1800.0, price=9.0, qty=200),    # 9円×200錠＝1,800円  → 合計3,600 → 税抜3,273 消費税327
    _pr('柏', 'k1', 500.0, price=5.0, qty=100),
]}
resv2 = [
    {'予約した店': '東立石', '出し手店': '氷川台', '対象年月': '202611', '受取予定月': '202701', '予約キー': 'h1', '薬品名': '薬h1', '予約日時': 'now'},
    {'予約した店': '東立石', '出し手店': '氷川台', '対象年月': '202611', '受取予定月': '202701', '予約キー': 'h2', '薬品名': '薬h2', '予約日時': 'now'},
    {'予約した店': '東立石', '出し手店': '氷川台', '対象年月': '202611', '受取予定月': '202701', '予約キー': 'gone', '薬品名': '消えた薬', '予約日時': 'now'},
    {'予約した店': '東立石', '出し手店': '柏', '対象年月': '202611', '受取予定月': '202611', '予約キー': 'k1', '薬品名': '薬k1', '予約日時': 'now'},
]
store_info = {'氷川台': {'法人名': '内観堂薬局', '薬局名（正式）': 'ダミー薬局 氷川台店', '所在地': 'ダミー県ダミー市1-1', '管理薬剤師名': 'ダミー 太郎'},
              '東立石': {'法人名': 'ソユーズ薬局', '薬局名（正式）': 'ダミー薬局 東立石店', '所在地': 'ダミー県ダミー市2-2', '管理薬剤師名': 'ダミー 花子'}}
data = al.build_pickup_request(result2, '東立石', resv2, '202611', store_info, sc.COMPANY_OF)
check('シート数＝出し手店数（2）', len(data['sheets']) == 2 and [s['出し手店'] for s in data['sheets']] == ['柏', '氷川台'])
sh_h = [s for s in data['sheets'] if s['出し手店'] == '氷川台'][0]
sh_k = [s for s in data['sheets'] if s['出し手店'] == '柏'][0]
check('氷川台→東立石は別法人（同一法人=False）／柏→東立石は同一法人', (not sh_h['同一法人']) and sh_k['同一法人'])
check('便＝2027年1月便（氷川台）／随時便（柏・今すぐ）', sh_h['便'] == '2027年1月便' and sh_k['便'] == '随時便（今すぐ）', (sh_h['便'], sh_k['便']))
check('単価・金額（データ側）', [(d['単価（薬価・税込）'], d['金額（税込）']) for d in sh_h['rows'] if d['状態'] == '出し手が掲載中'] == [(18.0, 1800.0), (9.0, 1800.0)])
check('一覧から外れた品も載る（単価・金額は空）', any(d['状態'].startswith('出し手の一覧から外れました') and d['単価（薬価・税込）'] == '' for d in sh_h['rows']))
bio = io.BytesIO(); yc.write_pickup_request_excel(bio, data); bio.seek(0)
wb = load_workbook(bio)
check('Excel シート名＝出し手店', wb.sheetnames == ['柏', '氷川台'])
ws = wb['氷川台']
# 明細ヘッダー行を探す
hdr_row = next(r for r in range(1, 40) if ws.cell(row=r, column=1).value == '薬品名')
headers = [ws.cell(row=hdr_row, column=c).value for c in range(1, 12)]
check('明細の列名が11列この順', headers == ['薬品名', '単位', '数量', '有効期限', 'ロットNO', '医薬品CD', '受取予定月', '区分', '単価（薬価・税込）', '金額（税込）', '状態'], headers)
labels = [ws.cell(row=r, column=1).value for r in range(3, 11)]
check('記名欄8行（便・発送日・出す店3行・受け取る店3行）', labels == ['便', '発送日', '出す店：法人名・薬局名', '出す店：所在地', '出す店：管理薬剤師名', '受け取る店：法人名・薬局名', '受け取る店：所在地', '受け取る店：管理薬剤師名'], labels)
vals = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=4).value for r in range(3, 11)}
check('記名欄の値（_店舗情報から）', vals['出す店：所在地'] == 'ダミー県ダミー市1-1' and vals['受け取る店：管理薬剤師名'] == 'ダミー 花子' and '内観堂薬局' in vals['出す店：法人名・薬局名'] and vals['便'] == '2027年1月便', vals)
r1 = hdr_row + 1
# 明細3行＝「消えた薬」（一覧から外れた品・薬品名順で先頭）・薬h1・薬h2。掲載中の行だけに式がある。
listed_rows = [r for r in range(r1, r1 + 3) if isinstance(ws.cell(row=r, column=9).value, (int, float))]
check('掲載中の2行の金額セルが =単価×数量 の式', len(listed_rows) == 2 and all(ws.cell(row=r, column=10).value == '=I%d*C%d' % (r, r) for r in listed_rows),
      [ws.cell(row=r, column=10).value for r in range(r1, r1 + 3)])
check('一覧から外れた行は式を置かない（空欄）', ws.cell(row=r1, column=10).value in (None, ''), ws.cell(row=r1, column=10).value)
# 合計欄（明細3行の下）
tot_row = hdr_row + 1 + 3
f_sum = ws.cell(row=tot_row, column=10).value
f_ex = ws.cell(row=tot_row + 1, column=10).value
f_tax = ws.cell(row=tot_row + 2, column=10).value
check('合計欄の式（SUM／ROUND(÷1.1,0)／差）', f_sum == '=SUM(J%d:J%d)' % (r1, r1 + 2) and f_ex == '=ROUND(J%d/1.1,0)' % tot_row and f_tax == '=J%d-J%d' % (tot_row, tot_row + 1), (f_sum, f_ex, f_tax))
# 式を解決した値（自前で計算）＝ 18×100 + 9×200 = 3,600 → 税抜3,273・消費税327
amt = sum((ws.cell(row=r, column=9).value or 0) * (float(str(ws.cell(row=r, column=3).value or 0).replace(',', '') or 0)) for r in range(r1, r1 + 3) if isinstance(ws.cell(row=r, column=9).value, (int, float)))
check('式の参照先を解決した値：税込3,600→税抜3,273・消費税327', amt == 3600 and round(amt / 1.1) == 3273 and amt - round(amt / 1.1) == 327, amt)
notes_h = [ws.cell(row=r, column=1).value for r in range(11, 15)]
check('別法人のシートに「同一法人」注記が無い', not any('同一法人内の移動' in str(v) for v in notes_h), notes_h)
wsk = wb['柏']
notes_k = [wsk.cell(row=r, column=1).value for r in range(11, 15)]
check('同一法人のシートに「同一法人内の移動（精算なし）」注記', any('同一法人内の移動（精算なし）' in str(v) for v in notes_k), notes_k)
kh = next(r for r in range(1, 40) if wsk.cell(row=r, column=1).value == '薬品名')
check('同一法人でも単価・金額は印字（式）', isinstance(wsk.cell(row=kh + 1, column=9).value, (int, float)) and str(wsk.cell(row=kh + 1, column=10).value).startswith('='))
total_w, budget_w, ok_w = yc.pickup_request_width_check()
check('列幅合計 %d ≤ 現行9列の合計 %d' % (total_w, budget_w), ok_w)
check('A4横・fitToWidth=1・見出し行の繰り返し', ws.page_setup.orientation == 'landscape' and ws.page_setup.fitToWidth == 1 and str(ws.print_title_rows).replace('$', '') == '%d:%d' % (hdr_row, hdr_row), (ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.print_title_rows))
check('記名欄が無い店は下線', '＿' in yc._store_field({}, '所在地'))

# ============================================================================
# 6) メール（smtplib を偽物に差し替え）
# ============================================================================
print('\n■ 6) mailer（偽物の SMTP）')

class FakeSMTP:
    instances = []
    fail_on = set()      # この宛先の送信で例外を出す
    def __init__(self, host, port, timeout=None, context=None):
        self.host, self.port, self.timeout = host, port, timeout
        self.logins = 0
        self.sent = []
        self.closed = False
        FakeSMTP.instances.append(self)
    def ehlo(self): pass
    def starttls(self, context=None): pass
    def login(self, u, p): self.logins += 1
    def send_message(self, msg):
        to = msg['To']
        if to in FakeSMTP.fail_on:
            FakeSMTP.fail_on.discard(to)
            raise RuntimeError('偽の送信失敗')
        self.sent.append((to, msg['Subject'], msg.get_content()))
    def quit(self): self.closed = True
    def close(self): self.closed = True

_orig_smtp, _orig_ssl = mailer.smtplib.SMTP, mailer.smtplib.SMTP_SSL
mailer.smtplib.SMTP = FakeSMTP
mailer.smtplib.SMTP_SSL = FakeSMTP
try:
    smtp_conf = {'host': 'smtp.example.invalid', 'port': 587, 'user': 'u', 'password': 'p', 'from': 'noreply@example.invalid'}
    emails36 = {n: 'dummy-%02d@example.invalid' % i for i, n in enumerate(sc.STORE_NAMES)}
    secrets36 = {'smtp': smtp_conf, 'app_url': 'https://app.example.invalid', 'store_emails': emails36}
    FakeSMTP.instances = []
    res = mailer.notify_allboard(secrets36, '東立石', '本文テスト', sc.STORE_NAMES)
    inst = FakeSMTP.instances
    check('全店板1投稿：接続1回', len(inst) == 1, len(inst))
    check('ログイン1回', inst and inst[0].logins == 1)
    check('35通（36店−自店）', len(res['sent']) == 35 and len(inst[0].sent) == 35, len(res['sent']))
    check('自店（東立石）には送らない', emails36['東立石'] not in res['sent'])
    def _store_of(addr): return next(n for n, a in emails36.items() if a == addr)
    ok_links = all(('store=' + mailer.quote(_store_of(to))) in body for to, _s, body in inst[0].sent)
    check('各本文に ?store=<その店> のリンク', ok_links)
    check('接続を閉じている', inst[0].closed)
    check('警告文言なし（全店宛先あり）', res['messages'] == [], res['messages'])
    # 14店だけ宛先あり → 13通＋要約1行
    emails14 = {n: emails36[n] for n in sc.SOYOUZ + sc.NAIKANDO}
    FakeSMTP.instances = []
    res = mailer.notify_allboard({'smtp': smtp_conf, 'app_url': 'https://app.example.invalid', 'store_emails': emails14}, '東立石', '本文', sc.STORE_NAMES)
    warn = [t for lv, t in res['messages'] if lv == 'warning']
    check('14店だけ→13通送信', len(res['sent']) == 13, len(res['sent']))
    check('「宛先未登録：22店（飛鳥22）」の要約1行', len(warn) == 1 and warn[0].startswith('宛先未登録：22店（飛鳥22）'), warn)
    # 途中1通が失敗しても残りは送られる（接続をつなぎ直して再送するので、その1通も届く）
    FakeSMTP.instances = []
    FakeSMTP.fail_on = {emails36['柏']}
    res = mailer.notify_allboard(secrets36, '東立石', '本文', sc.STORE_NAMES)
    check('途中1通が失敗→つなぎ直して再送し、35通そろう', len(res['sent']) == 35 and len(FakeSMTP.instances) == 2, (len(res['sent']), len(FakeSMTP.instances)))
    # 恒常的に失敗する宛先（再接続後も失敗）→ その1通だけ警告・残りは送られる
    class AlwaysFail(FakeSMTP):
        def send_message(self, msg):
            if msg['To'] == emails36['柏']:
                raise RuntimeError('偽の恒常失敗')
            self.sent.append((msg['To'], msg['Subject'], msg.get_content()))
    mailer.smtplib.SMTP = AlwaysFail
    FakeSMTP.instances = []
    res = mailer.notify_allboard(secrets36, '東立石', '本文', sc.STORE_NAMES)
    warn = [t for lv, t in res['messages'] if lv == 'warning']
    check('恒常失敗1通→34通送信＋その店の警告1行', len(res['sent']) == 34 and len(warn) == 1 and '柏' in warn[0], (len(res['sent']), warn))
    mailer.smtplib.SMTP = FakeSMTP
    # 予算切れ：1通ごとに時間がかかる偽物で budget=0.05秒 → 残りは送らず「○通は送れませんでした」
    class SlowSMTP(FakeSMTP):
        def send_message(self, msg):
            time.sleep(0.03)
            super().send_message(msg)
    mailer.smtplib.SMTP = SlowSMTP
    FakeSMTP.instances = []
    res = mailer.notify_allboard(secrets36, '東立石', '本文', sc.STORE_NAMES, budget=0.05)
    warn = [t for lv, t in res['messages'] if lv == 'warning']
    check('全体予算超過→残りを止め「○通は送れませんでした」が1行', res['skipped'] > 0 and len(warn) == 1 and ('%d通は送れませんでした' % res['skipped']) in warn[0] and len(res['sent']) + res['skipped'] == 35, (res['skipped'], warn))
    mailer.smtplib.SMTP = FakeSMTP
    # 1対1・予約通知：接続の timeout は 15秒のまま、宛先未登録は店名つき
    FakeSMTP.instances = []
    res = mailer.notify_new_message(secrets36, '東立石', '柏', '', '本文')
    check('1対1：timeout=15（不変）・1通', FakeSMTP.instances[0].timeout == 15 and len(res['sent']) == 1)
    res = mailer.notify_reservation({'smtp': smtp_conf, 'store_emails': {}}, '東立石', [{'出し手店': '柏', '薬品名': 'X'}])
    check('予約通知で宛先未登録→店名つきの警告（従来どおり）', any('柏' in t for lv, t in res['messages']))
    FakeSMTP.instances = []
    res = mailer.notify_allboard(secrets36, '東立石', '本文', sc.STORE_NAMES)
    check('全店板：timeout=8（不変）', FakeSMTP.instances[0].timeout == 8)
    check('[smtp] 未設定→未設定の案内1行・送信0', mailer.notify_allboard({}, '東立石', '本文', sc.STORE_NAMES)['messages'] == [('info', 'メール通知は未設定です。')])
finally:
    mailer.smtplib.SMTP, mailer.smtplib.SMTP_SSL = _orig_smtp, _orig_ssl

# ============================================================================
# 7) Googleシートの呼び出し回数（偽物のブック）
# ============================================================================
print('\n■ 7) gsheet_store（偽物のブック）')

class FakeWS:
    def __init__(self, title, grid):
        self.title = title
        self.grid = [list(r) for r in grid]
        self.updates = []
    def get_all_values(self):
        return [list(r) for r in self.grid]
    def update(self, range_name='A1', values=None, value_input_option='RAW'):
        self.updates.append((range_name, values))
        if range_name == 'A1' and values:
            self.grid = [list(r) for r in values]
    def clear(self):
        self.grid = []

class FakeBook:
    def __init__(self, sheets, batch_fail=False):
        self.id = 'fake-book-%d' % id(self)
        self._sheets = {ws.title: ws for ws in sheets}
        self.batch_calls = 0
        self.batch_fail = batch_fail
        self.raw_reads = 0
    def worksheets(self):
        return list(self._sheets.values())
    def add_worksheet(self, title, rows, cols):
        ws = FakeWS(title, [])
        self._sheets[title] = ws
        return ws
    def values_batch_get(self, ranges, params=None):
        self.batch_calls += 1
        if self.batch_fail:
            raise RuntimeError('偽の batchGet 失敗（タブが無い等）')
        out = []
        for rg in ranges:
            title = rg.strip("'")
            ws = self._sheets[title]
            out.append({'range': rg, 'values': ws.get_all_values()})
        return {'valueRanges': out}

header = list(yc.KEEP_COLS)
def _raw_ws(name):
    row = {c: '' for c in yc.KEEP_COLS}
    row.update({'個別医薬品CD': 'CD' + name, '薬品名': '薬' + name, '在庫数': '10', '薬価': '100', '薬価金額': '1000'})
    return FakeWS('raw_' + name, [header, [row[c] for c in header]])
index_grid = [gs.INDEX_HEADERS] + [[n, '202611', '2026/11/01 10:00', '1', 'OK', n + '_202611.csv'] for n in sc.STORE_NAMES]
book = FakeBook([FakeWS('_index', index_grid)] + [_raw_ws(n) for n in sc.STORE_NAMES])
gs.reset_ws_cache()
_orig_read_raw = gs.read_raw
def _count_read_raw(sh, name):
    sh.raw_reads += 1
    return _orig_read_raw(sh, name)
gs.read_raw = _count_read_raw
try:
    stores, latest, index = gs.load_current_month_stores(book)
    check('36店の raw 読み＝values_batch_get 1回・read_raw 0回', book.batch_calls == 1 and book.raw_reads == 0 and len(stores) == 36, (book.batch_calls, book.raw_reads, len(stores)))
    check('読んだ中身が正しい（薬品名・薬価）', stores[0]['rows'][0]['薬品名'] == '薬' + stores[0]['name'] and stores[0]['rows'][0]['薬価'] == '100')
    book2 = FakeBook([FakeWS('_index', index_grid)] + [_raw_ws(n) for n in sc.STORE_NAMES], batch_fail=True)
    gs.reset_ws_cache()
    stores2, _l, _i = gs.load_current_month_stores(book2)
    check('batchGet 失敗→従来の1タブずつ（read_raw 36回）に自動で落ちる', book2.batch_calls == 1 and book2.raw_reads == 36 and len(stores2) == 36, (book2.batch_calls, book2.raw_reads))
finally:
    gs.read_raw = _orig_read_raw
# _店舗情報 が無いブック → 空の辞書・例外なし・見出し行だけ作成
book3 = FakeBook([FakeWS('_index', index_grid)])
gs.reset_ws_cache()
gs._STORE_INFO_CACHE.clear()
info = gs.read_store_info(book3)
ws_info = book3._sheets.get(gs.STORE_INFO_TAB)
check('_店舗情報 が無い→空の辞書・例外なし', info == {})
check('見出し行だけ作成される', ws_info is not None and ws_info.get_all_values() == [gs.STORE_INFO_HEADERS], ws_info.get_all_values() if ws_info else None)
# あるブック → 列名で読む・60秒キャッシュ
book4 = FakeBook([FakeWS(gs.STORE_INFO_TAB, [gs.STORE_INFO_HEADERS, ['東立石', 'ソユーズ薬局', 'ダミー薬局 東立石店', 'ダミー県', 'ダミー 花子'], ['', 'x', 'x', 'x', 'x']])])
gs.reset_ws_cache(); gs._STORE_INFO_CACHE.clear()
info = gs.read_store_info(book4)
check('_店舗情報 を列名で読む（店名空の行は無視）', info == {'東立石': {'法人名': 'ソユーズ薬局', '薬局名（正式）': 'ダミー薬局 東立石店', '所在地': 'ダミー県', '管理薬剤師名': 'ダミー 花子'}}, info)
book4._sheets[gs.STORE_INFO_TAB].grid[1][3] = '変更後'
check('60秒キャッシュ（読み直さない）', gs.read_store_info(book4)['東立石']['所在地'] == 'ダミー県')
check('force=True で読み直す', gs.read_store_info(book4, force=True)['東立石']['所在地'] == '変更後')
lb = al.LocalBackend({})
check('LocalBackend.load_store_info は空辞書', lb.load_store_info() == {})

# ============================================================================
# 8) 見本・生成スクリプト・画面（2026-09-19 名前ごとの管理者）
# ============================================================================
print('\n■ 8) Secrets の見本・_検証用合言葉を作る.py・AppTest（名前つき管理者）')
import tomllib
import subprocess
_sample = os.path.join(HERE, '.streamlit', 'secrets.toml.sample')
try:
    _d = tomllib.loads(open(_sample, encoding='utf-8').read())
    check('secrets.toml.sample が TOML として読める', True)
    check('見本の [admin_passwords] は ADMIN_NAMES と同じ7人・同じ並び', list(_d.get('admin_passwords', {}).keys()) == sc.ADMIN_NAMES, list(_d.get('admin_passwords', {}).keys()))
    check('見本の [store_passwords] は STORE_NAMES と同じ36店・同じ並び', list(_d.get('store_passwords', {}).keys()) == sc.STORE_NAMES)
    _vals = list(_d['admin_passwords'].values()) + list(_d['store_passwords'].values())
    check('見本の値（管理者7＋店36）はすべて別（同じ値だと入れなくなるため）', len(set(_vals)) == len(_vals) == 43)
    check('見本の1行項目（spreadsheet_id・app_url）が表の中に紛れていない', 'spreadsheet_id' in _d and 'app_url' in _d and 'spreadsheet_id' not in _d['admin_passwords'])
    check('見本に旧 admin_password は有効行として残していない（コメント）', 'admin_password' not in _d)
except Exception as e:
    check('secrets.toml.sample が TOML として読める', False, e)
# 生成スクリプト（docs 側）… 引数なし・標準出力だけ・ファイルは作らない
_gen = os.path.normpath(os.path.join(HERE, '..', 'docs', '3法人対応改修_202610', '_検証用合言葉を作る.py'))
if os.path.exists(_gen):
    _before = set(os.listdir(os.path.dirname(_gen)))
    _p = subprocess.run([sys.executable, _gen], capture_output=True, text=True, cwd=os.path.dirname(_gen))
    _after = set(os.listdir(os.path.dirname(_gen)))
    try:
        _g = tomllib.loads(_p.stdout)
        check('_検証用合言葉を作る.py の出力が TOML として読める', _p.returncode == 0)
        check('出力は [admin_passwords] 7人（ADMIN_NAMES の順）＋[store_passwords] 36店', list(_g.get('admin_passwords', {}).keys()) == sc.ADMIN_NAMES and list(_g.get('store_passwords', {}).keys()) == sc.STORE_NAMES)
        _gv = list(_g['admin_passwords'].values()) + list(_g['store_passwords'].values())
        check('出力の43個の値がすべて別・空なし', len(set(_gv)) == 43 and all(_gv))
        check('旧 admin_password の1行は出さない', 'admin_password' not in _g)
        check('ファイルを作らない（フォルダの中身が増えていない）', _before == _after, _after - _before)
    except Exception as e:
        check('_検証用合言葉を作る.py の出力が TOML として読める', False, (e, _p.stderr[-300:]))
else:
    print('  [--] _検証用合言葉を作る.py が見つからないため省略')
# AppTest：ダミー Secrets を注入し、Googleシート未設定＝ローカル保管庫（LocalBackend）で画面を動かす
try:
    from streamlit.testing.v1 import AppTest

    def _ss(app, key):
        """ AppTest の session_state から安全に読む（.get が無く、無いキーは例外になるため）。 """
        try:
            return app.session_state[key]
        except Exception:
            return None
    _AP = {n: 'dummy-adm-%02d' % i for i, n in enumerate(sc.ADMIN_NAMES)}
    _SP = {n: 'dummy-%02d' % i for i, n in enumerate(sc.STORE_NAMES)}
    at = AppTest.from_file(os.path.join(HERE, 'streamlit_app.py'), default_timeout=120)
    at.secrets['admin_passwords'] = dict(_AP)
    at.secrets['store_passwords'] = dict(_SP)
    at.run()
    check('AppTest: 起動時に例外なし・合言葉の入力欄と「入る」がある', (not at.exception) and len(at.text_input) == 1 and any(b.label == '入る' for b in at.button), [str(e) for e in at.exception])
    at.text_input[0].input('wrong'); at.button[0].click().run()
    check('AppTest: 違う合言葉→エラー表示・入れない', (not _ss(at, 'authed')) and any('違います' in e.value for e in at.error))
    at.text_input[0].input(_AP['森田']); at.button[0].click().run()
    check('AppTest: 名前つき管理者（森田）で入れる（auth_mode=admin・auth_admin=森田）',
          _ss(at, 'authed') and _ss(at, 'auth_mode') == 'admin' and _ss(at, 'auth_admin') == '森田', 'session_state を表示できません')
    # ★2026-09-19 左バー→上の帯：本部モードの行・法人／店舗名の選択欄は本体（at.main）で探す。左バーには何も無い。
    check('AppTest: 上の帯（本体）に「本部モード：森田（36店を自由に切り替えできます）」',
          any(c.value == '本部モード：森田（36店を自由に切り替えできます）' for c in at.main.caption), [c.value for c in at.main.caption])
    check('AppTest: 左バーに要素が1つも無い（左バーをやめた・2026-09-19）',
          len(at.sidebar.children) == 0 and len(at.sidebar.caption) == 0 and len(at.sidebar.selectbox) == 0, len(at.sidebar.children))
    check('AppTest: 上の帯のアプリ名「💊 デッドストック」と「アップ済み」の数字が本体にある',
          any('デッドストック' in m.value for m in at.main.markdown) and any(m.label == 'アップ済み' for m in at.main.metric),
          ([m.value for m in at.main.markdown][:3], [m.label for m in at.main.metric]))
    _corp = [s for s in at.main.selectbox if s.label == '法人で絞る']
    _shop = [s for s in at.main.selectbox if s.label == '店舗名（必須）']
    check('AppTest: 「法人で絞る」（4択）と「店舗名」（36店＋先頭）が出る', len(_corp) == 1 and len(_corp[0].options) == 4 and len(_shop) == 1 and len(_shop[0].options) == 37,
          [(s.label, len(s.options)) for s in at.main.selectbox])
    check('AppTest: 店が未選択のあいだは赤字「まず店舗名を選んでください」が本体に出る',
          any('まず店舗名を選んでください' in m.value for m in at.main.markdown))
    _corp[0].select('飛鳥').run()
    _shop = [s for s in at.main.selectbox if s.label == '店舗名（必須）']
    check('AppTest: 飛鳥で絞ると店舗名は22店＋先頭', len(_shop) == 1 and len(_shop[0].options) == 23 and '本店' in _shop[0].options and '東立石' not in _shop[0].options)
    _shop[0].select('本店').run()
    check('AppTest: 本店を選ぶと my_store=本店・URL ?store=本店', _ss(at, 'my_store') == '本店' and at.query_params.get('store') in ('本店', ['本店']), (_ss(at, 'my_store'), dict(at.query_params)))
    check('AppTest: 店を選んだら赤字は消える', not any('まず店舗名を選んでください' in m.value for m in at.main.markdown))
    _corp = [s for s in at.main.selectbox if s.label == '法人で絞る'][0]
    _corp.select('内観堂').run()
    _shop = [s for s in at.main.selectbox if s.label == '店舗名（必須）'][0]
    _shop.select('氷川台').run()
    check('AppTest: 内観堂に切り替えて氷川台を選べる（管理者は店を自由に変えられる）', _ss(at, 'my_store') == '氷川台' and not at.exception, [str(e) for e in at.exception])
    check('AppTest: 店を切り替えても左バーは空のまま', len(at.sidebar.children) == 0)
    # 店の合言葉で入ると固定される（比較のため1本だけ）
    at2 = AppTest.from_file(os.path.join(HERE, 'streamlit_app.py'), default_timeout=120)
    at2.secrets['admin_passwords'] = dict(_AP)
    at2.secrets['store_passwords'] = dict(_SP)
    at2.run(); at2.text_input[0].input(_SP['東立石']); at2.button[0].click().run()
    check('AppTest: 店の合言葉で入ると東立石に固定・本部モードの行は出ない・法人の選択欄も出ない',
          _ss(at2, 'auth_store') == '東立石' and _ss(at2, 'auth_admin') is None
          and not any('本部モード' in c.value for c in at2.main.caption) and not any(s.label == '法人で絞る' for s in at2.main.selectbox))
    check('AppTest: 店の合言葉のときは上の帯に「自店：東立石（ソユーズ）」', any(m.value == '自店：**東立石**（ソユーズ）' for m in at2.main.markdown),
          [m.value for m in at2.main.markdown][:5])
    # 旧 admin_password だけの環境（[admin_passwords] なし）でも「本部」で入れる
    at3 = AppTest.from_file(os.path.join(HERE, 'streamlit_app.py'), default_timeout=120)
    at3.secrets['admin_password'] = 'dummy-admin-old'
    at3.secrets['store_passwords'] = dict(_SP)
    at3.run(); at3.text_input[0].input('dummy-admin-old'); at3.button[0].click().run()
    check('AppTest: 旧 admin_password だけでも入れて上の帯は「本部モード：本部（…）」',
          _ss(at3, 'auth_admin') == '本部' and any(c.value.startswith('本部モード：本部（') for c in at3.main.caption), [c.value for c in at3.main.caption])
except ImportError:
    print('  [--] streamlit.testing が無いため AppTest は省略')

print('\n合計：OK %d件 / NG %d件' % (_passed, _failed))
sys.exit(1 if _failed else 0)
